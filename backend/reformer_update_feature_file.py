import sys
import json
import os
import openpyxl

def main():
    try:
        # Load JSON configuration payload from stdin
        input_data = json.load(sys.stdin)
        
        raw_inputs = input_data.get('raw_inputs', [])
        inferred_tags = input_data.get('inferred_tags', [])
        kpis = input_data.get('kpis', [])
        plant_config = input_data.get('plant_config', {})
        
        # Resolve path to Easy_feature_file_reformer.xlsx relative to script directory
        script_dir = os.path.dirname(os.path.abspath(__file__))
        file_path = os.path.join(script_dir, 'Easy_feature_file_reformer.xlsx')
        
        # Fallback check
        if not os.path.exists(file_path):
            file_path = os.path.join(script_dir, 'easy_feature_file_reformer.xlsx')
            
        if not os.path.exists(file_path):
            file_path = os.path.join(os.getcwd(), 'server', 'Easy_feature_file_reformer.xlsx')
            if not os.path.exists(file_path):
                file_path = os.path.join(os.getcwd(), 'server', 'easy_feature_file_reformer.xlsx')
                
        if not os.path.exists(file_path):
            raise FileNotFoundError("Easy_feature_file_reformer.xlsx template file not found in server directory.")

        wb = openpyxl.load_workbook(file_path)
        sheet_names = wb.sheetnames
        
        pi_sheet_name = 'pi_tag_bank'
        tags_sheet_name = 'Tags'
        param_sheet_name = 'Pipeline_Parameters'
        
        for name in sheet_names:
            if name.lower() == 'pi_tag_bank':
                pi_sheet_name = name
            elif name.lower() == 'tags':
                tags_sheet_name = name
            elif name.lower() == 'pipeline_parameters':
                param_sheet_name = name
                
        # Resolve or create worksheets
        if pi_sheet_name in wb.sheetnames:
            ws_pi = wb[pi_sheet_name]
            if ws_pi.max_row > 1:
                ws_pi.delete_rows(2, ws_pi.max_row)
        else:
            ws_pi = wb.create_sheet(pi_sheet_name)
            
        if tags_sheet_name in wb.sheetnames:
            ws_tags = wb[tags_sheet_name]
            if ws_tags.max_row > 1:
                ws_tags.delete_rows(2, ws_tags.max_row)
        else:
            ws_tags = wb.create_sheet(tags_sheet_name)

        # Helper to convert to float/int if possible
        def to_num(val):
            if val is None or str(val).strip() == '' or str(val).strip() == '—':
                return None
            try:
                f = float(val)
                if f.is_integer():
                    return int(f)
                return f
            except ValueError:
                return val

        # Helper to extract physical tags from pi_name field
        def extract_tags(pi_name_val):
            if isinstance(pi_name_val, dict):
                return [t.strip() for t in pi_name_val.get('tags', []) if isinstance(t, str) and t.strip()]
            elif isinstance(pi_name_val, str):
                return [pi_name_val.strip()] if pi_name_val.strip() else []
            return []
            
        # Helper to format pi_name field for Tags sheet
        def get_formatted_pi_name(pi_name_val):
            if isinstance(pi_name_val, dict):
                method = pi_name_val.get('method') or 'Average'
                tags = [t.strip() for t in pi_name_val.get('tags', []) if isinstance(t, str) and t.strip()]
                if method == 'Custom':
                    formula = pi_name_val.get('formula') or ''
                    translated_formula = formula
                    for idx, t in enumerate(tags, 1):
                        placeholder = f"T{idx}"
                        translated_formula = translated_formula.replace(placeholder, t)
                    return f"Custom:{translated_formula}"
                else:
                    return f"{method}:{','.join(tags)}"
            return str(pi_name_val).strip() if pi_name_val else ""

        # 1. Overwrite pi_tag_bank header & rows
        headers_pi = ['model_name', 'PI_Name', 'Default_UOM', 'Min_value', 'max_value', 'default_value', 'CCP_rule']
        for col_idx, h in enumerate(headers_pi, 1):
            ws_pi.cell(row=1, column=col_idx, value=h)

        seen_pi_names = set()
        row_pi_idx = 2
        for tag in raw_inputs:
            pi_name_field = tag.get('pi_name')
            physical_tags = extract_tags(pi_name_field)
            
            for pi_name in physical_tags:
                if pi_name in seen_pi_names:
                    continue
                seen_pi_names.add(pi_name)
                
                ws_pi.cell(row=row_pi_idx, column=1, value="Live Benchmarking model")
                ws_pi.cell(row=row_pi_idx, column=2, value=pi_name)
                ws_pi.cell(row=row_pi_idx, column=3, value=tag.get('uom') or tag.get('default_uom') or '')
                ws_pi.cell(row=row_pi_idx, column=4, value=to_num(tag.get('min_value')))
                ws_pi.cell(row=row_pi_idx, column=5, value=to_num(tag.get('max_value')))
                ws_pi.cell(row=row_pi_idx, column=6, value=to_num(tag.get('default_value') or tag.get('design_value')))
                ws_pi.cell(row=row_pi_idx, column=7, value=None)
                row_pi_idx += 1

        # 2. Overwrite Tags header & rows
        headers_tags = [
            'model_name', 'name_short', 'pi_name', 'type', 'formula', 'hihi', 'lolo', 'high', 'low', 
            'uom', 'data_type', 'category', 'design_value', 'formula_state', 'pipeline_location', 
            'flag_benchmark', 'flag_actual', 'flag_clean', 'flag_ma', 'flag_actual_run_always', 
            'flag_inputcheck_nan', 'flag_dataclean_filter', 'flag_ma_filter', 'has_child', 'flag_clean_nan', 
            'ccp_lolo', 'ccp_hihi', 'ccp_default_value', 'ccp_oob_logic', 'ccp_stuck_logic', 
            'ccp_nan_logic', 'ccp_default_logic'
        ]
        for col_idx, h in enumerate(headers_tags, 1):
            ws_tags.cell(row=1, column=col_idx, value=h)

        seen_tags = set()
        row_tags_idx = 2
        
        # A. Append Raw PI input tags
        for tag in raw_inputs:
            name_short = tag.get('name_short') or ''
            if not name_short or name_short in seen_tags:
                continue
            seen_tags.add(name_short)
            
            formatted_pi_name = get_formatted_pi_name(tag.get('pi_name'))
            
            ws_tags.cell(row=row_tags_idx, column=1, value="Live Benchmarking model")
            ws_tags.cell(row=row_tags_idx, column=2, value=name_short)
            ws_tags.cell(row=row_tags_idx, column=3, value=formatted_pi_name)
            ws_tags.cell(row=row_tags_idx, column=4, value="pi")
            ws_tags.cell(row=row_tags_idx, column=5, value="")
            ws_tags.cell(row=row_tags_idx, column=10, value=tag.get('uom') or tag.get('default_uom') or '')
            ws_tags.cell(row=row_tags_idx, column=11, value="float")
            ws_tags.cell(row=row_tags_idx, column=12, value=tag.get('group') or '')
            ws_tags.cell(row=row_tags_idx, column=13, value=to_num(tag.get('design_value')))
            ws_tags.cell(row=row_tags_idx, column=14, value="active")
            ws_tags.cell(row=row_tags_idx, column=15, value="inline")
            ws_tags.cell(row=row_tags_idx, column=16, value=0)
            ws_tags.cell(row=row_tags_idx, column=17, value=1)
            ws_tags.cell(row=row_tags_idx, column=18, value=1)
            ws_tags.cell(row=row_tags_idx, column=19, value=1)
            ws_tags.cell(row=row_tags_idx, column=20, value=1)
            ws_tags.cell(row=row_tags_idx, column=21, value=1)
            ws_tags.cell(row=row_tags_idx, column=22, value=1)
            ws_tags.cell(row=row_tags_idx, column=23, value=1)
            ws_tags.cell(row=row_tags_idx, column=24, value=0)
            ws_tags.cell(row=row_tags_idx, column=25, value=1)
            ws_tags.cell(row=row_tags_idx, column=26, value=to_num(tag.get('min_value')))
            ws_tags.cell(row=row_tags_idx, column=27, value=to_num(tag.get('max_value')))
            ws_tags.cell(row=row_tags_idx, column=28, value=to_num(tag.get('default_value') or tag.get('design_value')))
            ws_tags.cell(row=row_tags_idx, column=29, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=30, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=31, value="default")
            ws_tags.cell(row=row_tags_idx, column=32, value="default")
            row_tags_idx += 1

        # B. Append Inferred tags
        for tag in inferred_tags:
            name_short = tag.get('name_short') or ''
            if not name_short or name_short in seen_tags:
                continue
            seen_tags.add(name_short)
            
            ws_tags.cell(row=row_tags_idx, column=1, value="Live Benchmarking model")
            ws_tags.cell(row=row_tags_idx, column=2, value=name_short)
            ws_tags.cell(row=row_tags_idx, column=3, value="")
            ws_tags.cell(row=row_tags_idx, column=4, value="inferred")
            ws_tags.cell(row=row_tags_idx, column=5, value=tag.get('formula') or '')
            ws_tags.cell(row=row_tags_idx, column=10, value=tag.get('uom') or '')
            ws_tags.cell(row=row_tags_idx, column=11, value="float")
            ws_tags.cell(row=row_tags_idx, column=12, value=tag.get('category') or '')
            ws_tags.cell(row=row_tags_idx, column=14, value="active")
            ws_tags.cell(row=row_tags_idx, column=15, value="inline")
            ws_tags.cell(row=row_tags_idx, column=16, value=0)
            ws_tags.cell(row=row_tags_idx, column=17, value=1)
            ws_tags.cell(row=row_tags_idx, column=18, value=1)
            ws_tags.cell(row=row_tags_idx, column=19, value=1)
            ws_tags.cell(row=row_tags_idx, column=20, value=1)
            ws_tags.cell(row=row_tags_idx, column=21, value=1)
            ws_tags.cell(row=row_tags_idx, column=22, value=1)
            ws_tags.cell(row=row_tags_idx, column=23, value=1)
            ws_tags.cell(row=row_tags_idx, column=24, value=0)
            ws_tags.cell(row=row_tags_idx, column=25, value=1)
            ws_tags.cell(row=row_tags_idx, column=29, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=30, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=31, value="default")
            ws_tags.cell(row=row_tags_idx, column=32, value="default")
            row_tags_idx += 1

        # C. Append KPIs
        for tag in kpis:
            name_short = tag.get('name_short') or ''
            if not name_short or name_short in seen_tags:
                continue
            seen_tags.add(name_short)
            
            tag_type = tag.get('type') or 'inferred'
            if tag_type != 'pi':
                tag_type = 'inferred'
                
            ws_tags.cell(row=row_tags_idx, column=1, value="Live Benchmarking model")
            ws_tags.cell(row=row_tags_idx, column=2, value=name_short)
            ws_tags.cell(row=row_tags_idx, column=3, value=(tag.get('pi_name') or '').strip())
            ws_tags.cell(row=row_tags_idx, column=4, value=tag_type)
            ws_tags.cell(row=row_tags_idx, column=5, value=tag.get('formula') or '')
            ws_tags.cell(row=row_tags_idx, column=10, value=tag.get('uom') or '')
            ws_tags.cell(row=row_tags_idx, column=11, value="float")
            ws_tags.cell(row=row_tags_idx, column=12, value=tag.get('category') or '')
            ws_tags.cell(row=row_tags_idx, column=14, value="active")
            ws_tags.cell(row=row_tags_idx, column=15, value="inline")
            ws_tags.cell(row=row_tags_idx, column=16, value=0)
            ws_tags.cell(row=row_tags_idx, column=17, value=1)
            ws_tags.cell(row=row_tags_idx, column=18, value=1)
            ws_tags.cell(row=row_tags_idx, column=19, value=1)
            ws_tags.cell(row=row_tags_idx, column=20, value=1)
            ws_tags.cell(row=row_tags_idx, column=21, value=1)
            ws_tags.cell(row=row_tags_idx, column=22, value=1)
            ws_tags.cell(row=row_tags_idx, column=23, value=1)
            ws_tags.cell(row=row_tags_idx, column=24, value=0)
            ws_tags.cell(row=row_tags_idx, column=25, value=1)
            ws_tags.cell(row=row_tags_idx, column=29, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=30, value="last_valid")
            ws_tags.cell(row=row_tags_idx, column=31, value="default")
            ws_tags.cell(row=row_tags_idx, column=32, value="default")
            row_tags_idx += 1

        # 3. Update Pipeline_Parameters with all plant config parameters
        if param_sheet_name in wb.sheetnames:
            ws_param = wb[param_sheet_name]
        else:
            ws_param = wb.create_sheet(param_sheet_name)
            
        # Clear existing parameters to rewrite clean
        if ws_param.max_row > 1:
            ws_param.delete_rows(2, ws_param.max_row)
        else:
            ws_param.cell(row=1, column=1, value='model_name')
            ws_param.cell(row=1, column=2, value='category')
            ws_param.cell(row=1, column=3, value='parameter')
            ws_param.cell(row=1, column=4, value='value')

        row_p_idx = 2
        for key, val in plant_config.items():
            ws_param.cell(row=row_p_idx, column=1, value="Live Benchmarking model")
            ws_param.cell(row=row_p_idx, column=2, value="Reformer_LBM")
            ws_param.cell(row=row_p_idx, column=3, value=key)
            ws_param.cell(row=row_p_idx, column=4, value=val)
            row_p_idx += 1

        # Save workbook in-place
        wb.save(file_path)
        wb.close()
        
        print(json.dumps({
            "status": "success", 
            "message": "Reformer Excel file updated successfully", 
            "file": file_path,
            "raw_inputs_count": len(seen_pi_names),
            "total_tags_count": len(seen_tags),
            "pi_sheet_written": pi_sheet_name,
            "tags_sheet_written": tags_sheet_name,
            "params_written": len(plant_config)
        }))
    except Exception as e:
        print(json.dumps({"status": "error", "message": str(e)}), file=sys.stderr)
        sys.exit(1)

if __name__ == '__main__':
    main()
